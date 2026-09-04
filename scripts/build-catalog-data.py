"""
Build a unified catalog JSON from the two source workbooks:

  - Zeovus_Life_14_Category_Revised_Catalogue.xlsx  -> nutraceuticals (14 sheets)
  - QUES_Product_Catalog_Descriptions.xlsx          -> cosmetics (5 sheets)

Output: scripts/catalog-seed.json

The output maps each source sheet to an EXISTING category slug (categories are
already seeded in the DB with icons/colours), so the seed only needs to wipe &
reinsert products. Products carry division-appropriate fields.
"""
import openpyxl, json, re, os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NUTRA_FILE = os.path.join(ROOT, 'Zeovus_Life_14_Category_Revised_Catalogue.xlsx')
COSMO_FILE = os.path.join(ROOT, 'QUES_Product_Catalog_Descriptions.xlsx')
OUT = os.path.join(ROOT, 'scripts', 'catalog-seed.json')

# Nutraceutical sheets are in the same order as DB sort_order 1..14.
NUTRA_SHEET_TO_SLUG = {
    '1. Healthy Ageing': 'healthy-ageing',
    '2. Multivitamins': 'multivitamins',
    '3. Gut Health': 'gut-health',
    "4. Women's Health": 'womens-health',
    "5. Men's Health": 'mens-health',
    '6. Brain Stress Sleep': 'brain-stress-sleep',
    '7. Immunity Respiratory': 'immunity',
    '8. Joint Bone Health': 'joint-bone',
    '9. Heart Health': 'heart-health',
    '10. Energy Sports': 'energy-sports',
    '11. Weight Mgmt': 'weight-management',
    '12. Beauty': 'beauty',
    "13. Children's Nutr": 'children',
    '14. Specialty Care': 'specialty',
}

COSMO_SHEET_TO_SLUG = {
    'Face Care': 'face-care',
    'Body Care': 'body-care',
    'Hair Care': 'hair-care',
    'Baby Care': 'baby-care',
    'Teen Care': 'teen-care',
}


def clean(v):
    if v is None:
        return None
    v = str(v).replace('\r\n', '\n').strip()
    return v if v else None


def slugify(s):
    s = re.sub(r'[^a-z0-9]+', '-', str(s).lower()).strip('-')
    return s[:145]


def find_header(rows, marker='Product Name'):
    for i, row in enumerate(rows[:8]):
        cells = [str(c).strip() if c is not None else '' for c in row]
        if marker in cells:
            return i, cells
    return 0, (rows[0] if rows else [])


def combine_benefits(b1, b2, b3):
    """Benefit 1 -> primary_benefit. Benefits 2 & 3 -> secondary_benefits (newline joined)."""
    primary = b1
    extras = [b for b in (b2, b3) if b]
    secondary = '\n'.join(extras) if extras else None
    return primary, secondary


def build_nutraceuticals():
    wb = openpyxl.load_workbook(NUTRA_FILE, data_only=True)
    categories = []
    for ws in wb.worksheets:
        slug = NUTRA_SHEET_TO_SLUG.get(ws.title)
        if not slug:
            print('  WARN: unmapped nutra sheet', ws.title)
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hi, headers = find_header(rows)
        col = {name: idx for idx, name in enumerate(headers)}

        def get(row, name):
            idx = col.get(name)
            return clean(row[idx]) if idx is not None and idx < len(row) else None

        products = []
        for row in rows[hi + 1:]:
            name = get(row, 'Product Name')
            if not name:
                continue
            b1 = get(row, 'Product Benefit 1')
            b2 = get(row, 'Product Benefit 2')
            b3 = get(row, 'Product Benefit 3')
            primary, secondary = combine_benefits(b1, b2, b3)
            products.append({
                'name': name,
                'subcat': None,
                'description': get(row, 'Description'),
                'key_actives': get(row, 'Key Actives'),
                'primary_benefit': primary,
                'secondary_benefits': secondary,
                'manufacturing_formats': get(row, 'Feasible Manufacturing Formats'),
                'dds_delivery_tech': get(row, 'Feasible Delivery Technology'),
                'recommended_dosage': get(row, 'Recommended Dosage'),
                'mechanism_of_action': get(row, 'Mechanism of Action'),
            })
        categories.append({'slug': slug, 'name': ws.title, 'division': 'nutraceuticals',
                           'subcategories': [], 'products': products})
    return categories


def build_cosmetics():
    wb = openpyxl.load_workbook(COSMO_FILE, data_only=True)
    categories = []
    for ws in wb.worksheets:
        slug = COSMO_SHEET_TO_SLUG.get(ws.title)
        if not slug:
            print('  WARN: unmapped cosmo sheet', ws.title)
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hi, headers = find_header(rows)
        col = {name: idx for idx, name in enumerate(headers)}

        def get(row, name):
            idx = col.get(name)
            return clean(row[idx]) if idx is not None and idx < len(row) else None

        subcats = {}
        sub_order = 0
        products = []
        for row in rows[hi + 1:]:
            name = get(row, 'Product Name')
            if not name:
                continue
            subcat = get(row, 'Sub-Category')
            if subcat and subcat not in subcats:
                sub_order += 1
                subcats[subcat] = sub_order
            products.append({
                'name': name,
                'subcat': subcat,
                'description': get(row, 'Description'),
                'skin_hair_type': get(row, 'Skin / Hair Type'),
                'concerns_addressed': get(row, 'Concerns Addressed'),
                'suitable_for': get(row, 'Suitable For'),
                'what_makes_potent': get(row, 'What Makes It Potent'),
                'manufacturing_formats': get(row, 'Standard Sizes Available'),
            })
        categories.append({
            'slug': slug, 'name': ws.title, 'division': 'cosmetics',
            'subcategories': [{'name': n, 'slug': slugify(n), 'sort_order': o} for n, o in subcats.items()],
            'products': products,
        })
    return categories


def main():
    print('Building nutraceuticals...')
    nutra = build_nutraceuticals()
    print('Building cosmetics...')
    cosmo = build_cosmetics()
    data = {'categories': nutra + cosmo}

    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    total = 0
    for c in data['categories']:
        total += len(c['products'])
        print(f"  [{len(c['products']):3}] {c['division']:14} {c['slug']:20} ({len(c['subcategories'])} subcats)")
    print(f"TOTAL PRODUCTS: {total}")
    print('Wrote', OUT)


if __name__ == '__main__':
    main()
