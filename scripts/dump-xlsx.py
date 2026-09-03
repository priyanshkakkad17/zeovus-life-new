import openpyxl, json
wb = openpyxl.load_workbook(r'C:\Users\DELL\Downloads\QUES_Product_Catalog_Descriptions.xlsx')
out = {}
for ws in wb.worksheets:
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    out[ws.title] = {
        "count": len(rows) - 1 if rows else 0,
        "headers": rows[0] if rows else [],
    }
print(json.dumps(out, indent=2, default=str))
with open(r'C:\Users\DELL\Desktop\zeovuslife\scripts\catalog-data.json','w',encoding='utf-8') as f:
    full = {ws.title: [list(r) for r in wb[ws.title].iter_rows(values_only=True)] for ws in wb.worksheets}
    json.dump(full, f, ensure_ascii=False, indent=2, default=str)
